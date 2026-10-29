"""Shared parsing helpers for tabular (CSV/XLS) bank statement providers.

Plain functions, not a class hierarchy: each helper is a stateless adapter
over one parsing seam (encoding detection, header search, cell decoding, …),
so there is no shared state to hang a base class off of, and a hierarchy
would add an interface layer without adding depth.
"""

from __future__ import annotations

import codecs
import csv
import io
import re
import shutil
import tempfile
import warnings
from collections.abc import Iterable, Sequence
from decimal import Decimal, InvalidOperation
from pathlib import Path

import xlrd
from openpyxl import load_workbook


def to_decimal(value: object) -> Decimal | None:
    """Convert a split-amount cell to Decimal.

    Returns None for an empty, placeholder, zero or unparsable cell, so callers
    can read "not this column" off the same result as "no usable value".
    """
    cleaned = str(value).replace(",", "").strip()
    if not cleaned or cleaned == "-":
        return None
    try:
        d = Decimal(cleaned)
        return d if d != 0 else None
    except InvalidOperation:
        return None


def normalize_cell_str(value: object) -> str:
    """Convert an xlrd cell value to text.

    xlrd hands back a float for every numeric cell, so an integral float is
    rendered without its ".0" tail to keep card suffixes and other digit strings
    comparable.
    """
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


# Ordered widest-first: the BOM variant must win before plain utf-8 swallows it.
CJK_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "gbk")


_BOM_ENCODINGS: tuple[tuple[bytes, str], ...] = (
    (codecs.BOM_UTF8, "utf-8-sig"),
    (codecs.BOM_UTF16_LE, "utf-16"),
    (codecs.BOM_UTF16_BE, "utf-16"),
)


def _bom_encoding(path: Path) -> str | None:
    with open(path, "rb") as f:
        head = f.read(3)
    for bom, encoding in _BOM_ENCODINGS:
        if head.startswith(bom):
            return encoding
    return None


def read_text_lines(path: Path, encodings: Sequence[str] = CJK_ENCODINGS) -> list[str]:
    """Read a text statement, trying each encoding in turn.

    A byte-order mark names the encoding outright, so it is tried before
    ``encodings``; some banks export UTF-16, which no CJK fallback can decode.
    """
    bom_encoding = _bom_encoding(path)
    if bom_encoding:
        encodings = (bom_encoding, *encodings)
    for encoding in encodings:
        try:
            with open(path, encoding=encoding) as f:
                return f.readlines()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError(f"Cannot decode {path}")


def find_header_line(
    lines: Sequence[str], keywords: Sequence[str], limit: int | None = None
) -> int | None:
    """Return the index of the first line containing every keyword."""
    for i, line in enumerate(lines[:limit]):
        if all(kw in line for kw in keywords):
            return i
    return None


CJK_HK_ENCODINGS: tuple[str, ...] = CJK_ENCODINGS + ("big5",)


def read_csv_rows(
    path: Path,
    *,
    encodings: Sequence[str] = CJK_HK_ENCODINGS,
    strip_cells: bool = True,
) -> list[list[str]]:
    """Read CSV rows, dropping rows whose every cell is blank.

    ``strip_cells`` also trims each cell, which banks that pad exports with tabs
    need before the header row can be matched.
    """
    for encoding in encodings:
        try:
            with open(path, encoding=encoding, newline="") as f:
                return [
                    [c.strip() for c in row] if strip_cells else row
                    for row in csv.reader(f)
                    if any(c.strip() for c in row)
                ]
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError(f"Cannot decode {path}")


MASKED_CARD_LAST4: re.Pattern[str] = re.compile(r"\d{4}\*+(\d{4})")


def masked_card_last4(text: str) -> str | None:
    """Extract the trailing 4 digits of a masked card number such as 6222****1234."""
    match = MASKED_CARD_LAST4.search(text)
    return match.group(1) if match else None


def first_masked_card_last4(lines: Iterable[str]) -> str | None:
    """Return the card suffix from the first line that carries a masked card number."""
    for line in lines:
        last4 = masked_card_last4(line)
        if last4:
            return last4
    return None


def join_description(*parts: str) -> str:
    """Join the non-empty description fragments; "Unknown" when none is present."""
    present = [p for p in parts if p]
    return " | ".join(present) if present else "Unknown"


def load_openpyxl_workbook(path: Path):
    """Load a workbook, tolerating banks that name an xlsx export ".xls".

    openpyxl refuses the .xls suffix outright, so such a file is copied to a
    temporary .xlsx before loading.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style",
            category=UserWarning,
        )
        if path.suffix.lower() == ".xls":
            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                    tmp_path = tmp.name
                    shutil.copy(path, tmp_path)
                return load_workbook(tmp_path)
            finally:
                if tmp_path:
                    Path(tmp_path).unlink(missing_ok=True)
        return load_workbook(path)


def _read_sheet_head(path: Path, nrows: int) -> list[list[str]]:
    data = path.read_bytes()
    # Banks ship OOXML under an .xls name, so the zip signature picks the reader.
    if data.startswith(b"PK"):
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style",
                category=UserWarning,
            )
            wb = load_workbook(io.BytesIO(data), read_only=True)
        try:
            sheet = wb.worksheets[0]
            # Read-only mode trusts the declared dimension, which some exports
            # understate, cutting rows down to their first cell.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_row=nrows, values_only=True)
            return [["" if v is None else str(v).strip() for v in row] for row in rows]
        finally:
            wb.close()
    sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
    return [
        [normalize_cell_str(v) for v in sheet.row_values(r)]
        for r in range(min(nrows, sheet.nrows))
    ]


def sheet_head_contains(path: Path, keywords: Iterable[str], nrows: int) -> bool:
    """Return whether every keyword is a substring of some cell in the first rows.

    Reads the first sheet of a BIFF or OOXML workbook regardless of its suffix,
    so renamed exports are recognised by content; an unreadable file is False.
    """
    try:
        cells = [cell for row in _read_sheet_head(path, nrows) for cell in row if cell]
    except Exception:
        return False
    return all(any(kw in cell for cell in cells) for kw in keywords)
